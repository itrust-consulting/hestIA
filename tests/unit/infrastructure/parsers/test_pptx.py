from __future__ import annotations

import io
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from hestia.domain.exceptions import ValidationError
from hestia.infrastructure.parsers.pptx import PPTXParser


def _make_pptx(tmp_path, slides=None):
    prs = Presentation()
    if slides is None:
        slides = [("Slide 1 Title", "Body text here")]
    for title_text, body_text in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title_text
        for ph in slide.placeholders:
            if ph.placeholder_format.idx == 1:
                ph.text = body_text
    path = tmp_path / "test.pptx"
    prs.save(str(path))
    return path


class TestPPTXParser:

    def test_loads_without_error(self, tmp_path):
        path = _make_pptx(tmp_path)
        p = PPTXParser(file=str(path))
        assert p.doc is not None

    def test_to_markdown_contains_title(self, tmp_path):
        path = _make_pptx(tmp_path, slides=[("My Slide", "Some content")])
        p = PPTXParser(file=str(path))
        md = p.to_markdown()
        assert "My Slide" in md

    def test_to_markdown_contains_body(self, tmp_path):
        path = _make_pptx(tmp_path, slides=[("Title", "Important body text")])
        p = PPTXParser(file=str(path))
        md = p.to_markdown()
        assert "Important body text" in md

    def test_multiple_slides(self, tmp_path):
        path = _make_pptx(tmp_path, slides=[
            ("Slide A", "Content A"),
            ("Slide B", "Content B"),
        ])
        p = PPTXParser(file=str(path))
        md = p.to_markdown()
        assert "Slide A" in md
        assert "Slide B" in md

    def test_metadata_has_source(self, tmp_path):
        path = _make_pptx(tmp_path)
        p = PPTXParser(file=str(path))
        meta = p.get_metadata()
        assert "source" in meta

    def test_get_metadata_applies_mask(self, tmp_path):
        path = _make_pptx(tmp_path)
        p = PPTXParser(file=str(path))
        meta = p.get_metadata(mask_name="rag_default")
        assert "source" in meta

    def test_to_markdown_skips_non_text_shapes_empty_paragraphs_and_empty_slides(self, tmp_path):
        prs = Presentation()
        # Slide 1: has a table shape (no text frame) and an empty body placeholder
        # (empty paragraph) alongside the title -> only the heading survives,
        # so the whole slide is dropped from the output.
        slide1 = prs.slides.add_slide(prs.slide_layouts[1])
        slide1.shapes.title.text = "Empty Slide"
        slide1.shapes.add_table(2, 2, Pt(0), Pt(0), Pt(100), Pt(100))

        # Slide 2: has real body content so it should appear in the output.
        slide2 = prs.slides.add_slide(prs.slide_layouts[1])
        slide2.shapes.title.text = "Real Slide"
        for ph in slide2.placeholders:
            if ph.placeholder_format.idx == 1:
                ph.text = "Actual content"

        path = tmp_path / "mixed.pptx"
        prs.save(str(path))

        p = PPTXParser(file=str(path))
        md = p.to_markdown()
        assert "Empty Slide" not in md
        assert "Real Slide" in md
        assert "Actual content" in md


class TestPPTXParserMetadataCoreProperties:
    """Exercise get_metadata's core-properties branches with a fake doc,
    since python-pptx's default CoreProperties always has a truthy
    `modified` and no attribute matching a falsy/absent case for every
    field at once."""

    class _FakeCoreProps:
        def __init__(self, **kwargs):
            self.title = kwargs.get("title")
            self.subject = kwargs.get("subject")
            self.author = kwargs.get("author")
            self.version = kwargs.get("version")
            self.language = kwargs.get("language")
            self.modified = kwargs.get("modified")

    class _FakeDoc:
        def __init__(self, core_properties):
            self.core_properties = core_properties

    def test_all_props_falsy_and_no_modified_adds_nothing(self):
        p = PPTXParser.__new__(PPTXParser)
        p.filepath = "some/file.pptx"
        p.doc = self._FakeDoc(self._FakeCoreProps())
        meta = p.get_metadata()
        assert "title" not in meta
        assert "modified" not in meta
        assert meta["source"] == "file"

    def test_truthy_props_and_modified_are_extracted(self):
        import datetime
        p = PPTXParser.__new__(PPTXParser)
        p.filepath = "some/file.pptx"
        p.doc = self._FakeDoc(self._FakeCoreProps(
            title="My Title", modified=datetime.datetime(2024, 1, 1)
        ))
        meta = p.get_metadata()
        assert meta["title"] == "My Title"
        assert meta["modified"] == "2024-01-01T00:00:00"

    def test_doc_none_skips_core_properties_entirely(self):
        p = PPTXParser.__new__(PPTXParser)
        p.filepath = "some/file.pptx"
        p.doc = None
        meta = p.get_metadata()
        assert meta == {"source": "file", "source_uri": "some/file.pptx"}


class TestPPTXParserEmptyDoc:

    def test_to_markdown_raises_when_doc_is_none(self):
        p = PPTXParser.__new__(PPTXParser)
        p.doc = None
        with pytest.raises(ValidationError, match="Empty document"):
            p.to_markdown()


class TestSlideTitleHelper:

    def test_returns_title_text(self, tmp_path):
        path = _make_pptx(tmp_path, slides=[("My Title", "Body")])
        p = PPTXParser(file=str(path))
        slide = p.doc.slides[0]
        assert p._slide_title(slide) == "My Title"

    def test_returns_empty_string_for_untitled_slide(self, tmp_path):
        from pptx import Presentation
        prs = Presentation()
        # Use a blank layout (no title placeholder)
        blank_layout = prs.slide_layouts[6]  # "Blank" layout
        prs.slides.add_slide(blank_layout)
        path = tmp_path / "blank.pptx"
        prs.save(str(path))

        p = PPTXParser(file=str(path))
        slide = p.doc.slides[0]
        result = p._slide_title(slide)
        assert result == ""
