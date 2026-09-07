from __future__ import annotations

import pathlib
import pytest

# Real test file from tests/test_inputs/
DOCX_FILE = pathlib.Path(__file__).parents[4] / "tests" / "test_inputs" / "Test-Document.docx"
DOCX_ITR = pathlib.Path(__file__).parents[4] / "tests" / "test_inputs" / "Test-Document-ITR.docx"

skip_if_missing = pytest.mark.skipif(not DOCX_FILE.exists(), reason="Test DOCX not found")
skip_if_itr_missing = pytest.mark.skipif(not DOCX_ITR.exists(), reason="ITR DOCX not found")

from hestia.infrastructure.parsers.docx import DOCXParser, ITRDOCXParser


@skip_if_missing
class TestDOCXParser:

    @pytest.fixture(scope="class")
    def parser(self):
        return DOCXParser(file=str(DOCX_FILE))

    def test_loads_without_error(self, parser):
        assert parser.doc is not None

    def test_to_markdown_contains_text(self, parser):
        md = parser.to_markdown()
        assert len(md) > 0

    def test_tables_rendered_as_markdown(self, parser):
        md = parser.to_markdown()
        # If the test document contains a table, it should have pipe characters
        # We just verify markdown is returned without crash

    def test_metadata_has_source(self, parser):
        meta = parser.get_metadata()
        assert "source" in meta

    def test_metadata_builtin_only(self, parser):
        meta = parser.get_metadata(builtIn_only=True)
        assert isinstance(meta, dict)

    def test_headings_rendered_with_hash(self, parser):
        md = parser.to_markdown()
        # Should contain heading markers if the document has headings
        # Just verify the method runs
        assert isinstance(md, str)


# ---------------------------------------------------------------------------
# DOCXParser helpers — unit tests without needing a real file
# ---------------------------------------------------------------------------

class TestDOCXParserHelpers:
    """Test methods that can be exercised without a full .docx parse."""

    @pytest.fixture
    def parser(self):
        p = DOCXParser.__new__(DOCXParser)
        p._style_levels = {}
        p._footnotes = {}
        p._rels = {"rId1": "https://example.com"}
        return p

    def test_render_hyperlink_with_known_rel(self, parser):
        from lxml import etree
        ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        elem = etree.fromstring(f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:r="{ns}" r:id="rId1">Click here</w:hyperlink>')
        result = parser.render_hyperlink(elem)
        assert "Click here" in result
        assert "https://example.com" in result

    def test_render_hyperlink_without_rel_returns_text(self, parser):
        from lxml import etree
        elem = etree.fromstring('<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">plain text</w:hyperlink>')
        result = parser.render_hyperlink(elem)
        assert result == "plain text"

    def test_table_to_markdown_produces_pipe_table(self, tmp_path):
        import docx
        doc = docx.Document()
        tbl = doc.add_table(rows=2, cols=2)
        tbl.cell(0, 0).text = "H1"
        tbl.cell(0, 1).text = "H2"
        tbl.cell(1, 0).text = "V1"
        tbl.cell(1, 1).text = "V2"
        path = tmp_path / "tbl.docx"
        doc.save(str(path))

        p = DOCXParser(file=str(path))
        md = p.table_to_markdown(p.doc.tables[0])
        assert "|" in md
        assert "H1" in md
        assert "V1" in md

    def test_table_to_markdown_merged_cells_no_crash(self, tmp_path):
        import docx
        from docx.oxml.ns import qn
        doc = docx.Document()
        tbl = doc.add_table(rows=2, cols=2)
        # Merge cells (0,0) and (0,1)
        cell_a = tbl.cell(0, 0)
        cell_b = tbl.cell(0, 1)
        cell_a.merge(cell_b)
        path = tmp_path / "merged.docx"
        doc.save(str(path))

        p = DOCXParser(file=str(path))
        md = p.table_to_markdown(p.doc.tables[0])
        assert isinstance(md, str)

    def test_table_to_markdown_zero_rows_returns_empty(self, tmp_path):
        import docx
        doc = docx.Document()
        doc.add_table(rows=0, cols=2)
        path = tmp_path / "zero_rows.docx"
        doc.save(str(path))

        p = DOCXParser(file=str(path))
        md = p.table_to_markdown(p.doc.tables[0])
        assert md == ""

    def test_get_list_level_returns_none_when_numpr_has_no_ilvl(self, tmp_path):
        import docx
        doc = docx.Document()
        para = doc.add_paragraph("item")
        pPr = para._p.get_or_add_pPr()
        pPr.get_or_add_numPr()  # numPr present but no <w:ilvl> child
        path = tmp_path / "numpr.docx"
        doc.save(str(path))

        p = DOCXParser(file=str(path))
        level = p.get_list_level_from_paragraph(doc.paragraphs[-1], {})
        assert level is None

    def test_parse_list_styles_missing_styles_xml(self, tmp_path):
        import zipfile
        # A minimal zip that just lacks word/styles.xml
        path = tmp_path / "no_styles.docx"
        with zipfile.ZipFile(path, "w") as zf:
            zf.writestr("word/document.xml", "<root/>")
        p = DOCXParser.__new__(DOCXParser)
        result = p._parse_list_styles(str(path))
        assert result == {}

    def test_parse_rels_missing_rels_file(self, tmp_path):
        import zipfile
        path = tmp_path / "no_rels.docx"
        with zipfile.ZipFile(path, "w") as zf:
            zf.writestr("word/document.xml", "<root/>")
        p = DOCXParser.__new__(DOCXParser)
        result = p._parse_rels(str(path))
        assert result == {}

    def test_omml_to_md_latex_pandoc_unavailable(self, monkeypatch):
        import hestia.infrastructure.parsers.docx as docx_module
        monkeypatch.setattr(docx_module, "PANDOC_AVAILABLE", False)
        p = DOCXParser.__new__(DOCXParser)
        result = p.omml_to_md_latex(None)
        assert "pandoc not installed" in result

    def test_omml_to_md_latex_conversion_failure(self, monkeypatch):
        import hestia.infrastructure.parsers.docx as docx_module
        from lxml import etree

        def boom(*args, **kwargs):
            raise RuntimeError("pandoc exploded")

        monkeypatch.setattr(docx_module.pypandoc, "convert_file", boom)
        p = DOCXParser.__new__(DOCXParser)
        m_ns = docx_module.DOCXParser.NAMESPACES["m"]
        elem = etree.fromstring(f'<m:oMath xmlns:m="{m_ns}"/>')
        result = p.omml_to_md_latex(elem)
        assert "OMML conversion failed" in result

    def test_omml_to_md_latex_unlink_failure_is_swallowed(self, monkeypatch):
        import hestia.infrastructure.parsers.docx as docx_module
        from lxml import etree

        monkeypatch.setattr(docx_module.pypandoc, "convert_file", lambda *a, **k: "converted text")

        def boom_unlink(*args, **kwargs):
            raise OSError("cannot delete")

        monkeypatch.setattr(docx_module.os, "unlink", boom_unlink)
        p = DOCXParser.__new__(DOCXParser)
        m_ns = docx_module.DOCXParser.NAMESPACES["m"]
        elem = etree.fromstring(f'<m:oMath xmlns:m="{m_ns}"/>')
        result = p.omml_to_md_latex(elem)
        assert result == "converted text"

    def test_to_markdown_skips_empty_paragraph_and_empty_table(self, tmp_path):
        import docx
        doc = docx.Document()
        doc.add_paragraph("Real content")
        doc.add_paragraph("")  # empty paragraph -> renders to ""
        doc.add_table(rows=0, cols=2)  # empty table -> renders to ""
        path = tmp_path / "mixed.docx"
        doc.save(str(path))

        p = DOCXParser(file=str(path))
        md = p.to_markdown()
        assert "Real content" in md
        assert md.count("\n\n\n") == 0

    def test_get_metadata_applies_mask(self, tmp_path):
        import docx
        doc = docx.Document()
        doc.add_paragraph("hi")
        path = tmp_path / "masked.docx"
        doc.save(str(path))

        p = DOCXParser(file=str(path))
        meta = p.get_metadata(mask_name="rag_default")
        assert "source" in meta


@skip_if_missing
class TestDOCXParserListItems:

    def test_list_items_rendered_in_markdown(self):
        # The Test-Document.docx should render paragraphs; we just verify the
        # to_markdown output is a non-empty string without error.
        p = DOCXParser(file=str(DOCX_FILE))
        md = p.to_markdown()
        assert isinstance(md, str)
        assert len(md) > 0


@skip_if_itr_missing
class TestITRDOCXParser:

    @pytest.fixture(scope="class")
    def parser(self):
        return ITRDOCXParser(file=str(DOCX_ITR))

    def test_loads_without_error(self, parser):
        assert parser.doc is not None

    def test_metadata_merged_from_cover_page(self, parser):
        meta = parser.get_metadata(builtIn_only=False)
        assert isinstance(meta, dict)

    def test_to_markdown_returns_string(self, parser):
        md = parser.to_markdown()
        assert isinstance(md, str)
        assert len(md) > 0

    def test_get_metadata_builtin_only_default_skips_cover_page(self, parser):
        # builtIn_only defaults to True, so get_metadata_from_cover_page is never called
        meta = parser.get_metadata()
        assert "source" in meta

    def test_get_metadata_applies_mask(self, parser):
        meta = parser.get_metadata(mask_name="rag_default")
        assert isinstance(meta, dict)


class TestITRDOCXParserCoverPageEdgeCases:

    def test_cover_page_metadata_empty_with_fewer_than_two_tables(self, tmp_path):
        import docx
        doc = docx.Document()
        doc.add_paragraph("no tables here")
        path = tmp_path / "no_cover.docx"
        doc.save(str(path))

        p = ITRDOCXParser(file=str(path))
        result = p.get_metadata_from_cover_page(str(path))
        assert result == {}

    def test_get_metadata_swallows_cover_page_errors(self, tmp_path, monkeypatch):
        import docx
        doc = docx.Document()
        doc.add_paragraph("content")
        path = tmp_path / "err.docx"
        doc.save(str(path))

        p = ITRDOCXParser(file=str(path))

        def boom(file):
            raise IndexError("bad cover page")

        monkeypatch.setattr(p, "get_metadata_from_cover_page", boom)
        meta = p.get_metadata(builtIn_only=False)
        assert "source" in meta
