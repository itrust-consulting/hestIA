from __future__ import annotations

import openpyxl
import pytest

from hestia.infrastructure.parsers.xlsx import XLSXParser


def _make_xlsx(tmp_path, sheet_data=None):
    if sheet_data is None:
        sheet_data = {"Sheet1": [["Name", "Age"], ["Alice", 30], ["Bob", 25]]}
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet_name, rows in sheet_data.items():
        ws = wb.create_sheet(sheet_name)
        for row in rows:
            ws.append(row)
    path = tmp_path / "test.xlsx"
    wb.save(str(path))
    return path


class TestXLSXParser:

    def test_loads_without_error(self, tmp_path):
        path = _make_xlsx(tmp_path)
        p = XLSXParser(file=str(path))
        assert p.doc is not None

    def test_loads_all_sheets(self, tmp_path):
        path = _make_xlsx(tmp_path, sheet_data={
            "Sheet1": [["A", "B"], [1, 2]],
            "Sheet2": [["X", "Y"], [3, 4]],
        })
        p = XLSXParser(file=str(path))
        assert "Sheet1" in p.doc
        assert "Sheet2" in p.doc

    def test_selected_sheets_filter(self, tmp_path):
        path = _make_xlsx(tmp_path, sheet_data={
            "Sheet1": [["A"], [1]],
            "Sheet2": [["B"], [2]],
        })
        p = XLSXParser(file=str(path), selected_sheets=["Sheet1"])
        assert "Sheet1" in p.doc
        assert "Sheet2" not in p.doc

    def test_to_markdown_contains_data(self, tmp_path):
        path = _make_xlsx(tmp_path)
        p = XLSXParser(file=str(path))
        md = p.to_markdown()
        assert "Alice" in md
        assert "Bob" in md

    def test_metadata_has_source(self, tmp_path):
        path = _make_xlsx(tmp_path)
        p = XLSXParser(file=str(path))
        meta = p.get_metadata()
        assert "source" in meta

    def test_to_markdown_sheets_returns_dict(self, tmp_path):
        path = _make_xlsx(tmp_path, sheet_data={
            "S1": [["A"], [1]],
            "S2": [["B"], [2]],
        })
        p = XLSXParser(file=str(path))
        sheets = p.to_markdown_sheets()
        assert "S1" in sheets
        assert "S2" in sheets
        assert isinstance(sheets["S1"], str)


# ---------------------------------------------------------------------------
# XLSXParser.to_csv
# ---------------------------------------------------------------------------

class TestXLSXToCSV:

    def test_to_csv_returns_csv_string(self, tmp_path):
        path = _make_xlsx(tmp_path, sheet_data={"Data": [["Col1", "Col2"], ["a", "b"]]})
        p = XLSXParser(file=str(path))
        result = p.to_csv()
        assert "Data" in result
        csv_content = result["Data"]
        assert "Col1" in csv_content or "0" in csv_content  # header row

    def test_to_csv_writes_file_when_output_path_given(self, tmp_path):
        path = _make_xlsx(tmp_path, sheet_data={"Sheet1": [["A", "B"], [1, 2]]})
        p = XLSXParser(file=str(path))
        out_dir = tmp_path / "output"
        out_dir.mkdir()
        result = p.to_csv(output_path=str(out_dir))
        assert "Sheet1" in result
        import os
        assert os.path.exists(result["Sheet1"])


# ---------------------------------------------------------------------------
# XLSXParser.get_builtin_metadata
# ---------------------------------------------------------------------------

class TestXLSXBuiltinMetadata:

    def test_get_builtin_metadata_returns_dict(self, tmp_path):
        path = _make_xlsx(tmp_path)
        p = XLSXParser(file=str(path))
        meta = p.get_builtin_metadata()
        assert isinstance(meta, dict)

    def test_get_builtin_metadata_includes_custom_properties(self, tmp_path):
        from openpyxl.packaging.custom import StringProperty
        wb = openpyxl.Workbook()
        wb.active.append(["A", "B"])
        wb.custom_doc_props.append(StringProperty(name="MyProp", value="MyVal"))
        path = tmp_path / "custom.xlsx"
        wb.save(str(path))

        p = XLSXParser(file=str(path))
        meta = p.get_builtin_metadata()
        assert meta.get("myprop") == "MyVal"

    def test_get_metadata_builtin_only_false_uses_empty_cover_page(self, tmp_path):
        # Base XLSXParser.get_metadata_from_cover_page always returns {}.
        path = _make_xlsx(tmp_path)
        p = XLSXParser(file=str(path))
        meta = p.get_metadata(builtIn_only=False)
        assert "source" in meta

    def test_get_metadata_swallows_cover_page_errors(self, tmp_path, monkeypatch):
        path = _make_xlsx(tmp_path)
        p = XLSXParser(file=str(path))

        def boom():
            raise AttributeError("cover page exploded")

        monkeypatch.setattr(p, "get_metadata_from_cover_page", boom)
        meta = p.get_metadata(builtIn_only=False)
        assert "source" in meta

    def test_get_metadata_from_cover_page_base_returns_empty(self, tmp_path):
        path = _make_xlsx(tmp_path)
        p = XLSXParser(file=str(path))
        assert p.get_metadata_from_cover_page() == {}

    def test_get_metadata_applies_mask(self, tmp_path):
        path = _make_xlsx(tmp_path)
        p = XLSXParser(file=str(path))
        meta = p.get_metadata(mask_name="rag_default")
        assert "source" in meta


class TestXLSXParserEmptyDocErrors:

    def test_to_markdown_sheets_raises_when_doc_is_none(self):
        from hestia.domain.exceptions import ValidationError
        p = XLSXParser.__new__(XLSXParser)
        p.doc = None
        with pytest.raises(ValidationError, match="Empty document"):
            p.to_markdown_sheets()

    def test_to_csv_raises_when_doc_is_none(self):
        from hestia.domain.exceptions import ValidationError
        p = XLSXParser.__new__(XLSXParser)
        p.doc = None
        with pytest.raises(ValidationError, match="Empty document"):
            p.to_csv()


class TestXLSXToCSVSheetIds:

    def test_to_csv_with_single_int_sheet_id(self, tmp_path):
        path = _make_xlsx(tmp_path, sheet_data={
            "Sheet1": [["A"], [1]],
            "Sheet2": [["B"], [2]],
        })
        p = XLSXParser(file=str(path))
        result = p.to_csv(sheet_ids=1)
        assert list(result.keys()) == ["Sheet2"]

    def test_to_csv_with_list_of_sheet_ids(self, tmp_path):
        path = _make_xlsx(tmp_path, sheet_data={
            "Sheet1": [["A"], [1]],
            "Sheet2": [["B"], [2]],
            "Sheet3": [["C"], [3]],
        })
        p = XLSXParser(file=str(path))
        result = p.to_csv(sheet_ids=[0, 2])
        assert set(result.keys()) == {"Sheet1", "Sheet3"}
