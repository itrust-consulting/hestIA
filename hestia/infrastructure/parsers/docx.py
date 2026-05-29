from __future__ import annotations

import logging
import os
import zipfile
from pathlib import Path
from tempfile import NamedTemporaryFile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree

from hestia.infrastructure.parsers.base import BaseParser, MetadataFilter

try:
    import pypandoc
    PANDOC_AVAILABLE = True
except ModuleNotFoundError:
    PANDOC_AVAILABLE = False

_log = logging.getLogger("hestia.system")


class DOCXParser(BaseParser):

    NAMESPACES = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "dcelem": "http://purl.org/dc/elements/1.1/",
        "dcterm": "http://purl.org/dc/terms/",
        "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
        "ep": "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties",
        "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
        "dgm": "http://schemas.openxmlformats.org/package/2006/relationships",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }

    def __init__(self, file=None):
        self._style_levels = None
        self._footnotes: dict = {}
        self._rels: dict = {}
        self._loader_map = {".docx": self._load_doc}
        super().__init__(file=file)

    def _load_doc(self, file):
        self._style_levels = self._parse_list_styles(file)
        self._footnotes = self._parse_footnotes(file)
        self._rels = self._parse_rels(file)
        _log.debug("docx_loaded", extra={"file": str(file), "n_rels": len(self._rels), "n_footnotes": len(self._footnotes)})
        return Document(file)

    # ------------------------------------------------------------------
    # XML extraction helpers
    # ------------------------------------------------------------------

    def _parse_list_styles(self, file) -> dict:
        style_levels = {}
        with zipfile.ZipFile(file) as zf:
            if "word/styles.xml" not in zf.namelist():
                return style_levels
            root = etree.fromstring(zf.read("word/styles.xml"))
        for style in root.findall("w:style", self.NAMESPACES):
            style_id = style.get("{%s}styleId" % self.NAMESPACES["w"])
            numpr = style.find(".//w:numPr", self.NAMESPACES)
            if numpr is not None:
                ilvl_elem = numpr.find("w:ilvl", self.NAMESPACES)
                ilvl = int(ilvl_elem.get("{%s}val" % self.NAMESPACES["w"])) if ilvl_elem is not None else 0
                style_levels[style_id] = ilvl
        return style_levels

    def _parse_footnotes(self, file) -> dict:
        footnotes = {}
        with zipfile.ZipFile(file) as zf:
            if "word/footnotes.xml" not in zf.namelist():
                return footnotes
            root = etree.fromstring(zf.read("word/footnotes.xml"))
        for footnote in root.findall("w:footnote", self.NAMESPACES):
            ref = footnote.get("{%s}id" % self.NAMESPACES["w"])
            text = [t.text for t in footnote.findall(".//w:t", self.NAMESPACES) if t.text]
            footnotes[ref] = "".join(text)
        return footnotes

    def _parse_rels(self, file) -> dict:
        rels = {}
        with zipfile.ZipFile(file) as zf:
            if "word/_rels/document.xml.rels" not in zf.namelist():
                return rels
            root = etree.fromstring(zf.read("word/_rels/document.xml.rels"))
        for rel in root.iter("{%s}Relationship" % self.NAMESPACES["dgm"]):
            rels[rel.get("Id")] = rel.get("Target")
        return rels

    # ------------------------------------------------------------------
    # Rendering helpers
    # ------------------------------------------------------------------

    def get_list_level_from_paragraph(self, para: Paragraph, style_levels: dict) -> int | None:
        pPr = para._p.pPr
        if pPr is None:
            return None
        if pPr.numPr is not None:
            ilvl = pPr.numPr.ilvl
            if ilvl is not None:
                return ilvl.val
        if pPr.pStyle is not None:
            style_id = pPr.pStyle.val
            if style_id in style_levels:
                return style_levels[style_id]
        return None

    def omml_to_md_latex(self, omml_elem) -> str:
        if not PANDOC_AVAILABLE:
            _log.warning("omml_conversion_skipped", extra={"reason": "pypandoc not installed"})
            return "<--OMML conversion failed: pandoc not installed.-->"

        tmp_path = None
        try:
            tmp_doc = Document()
            tmp_doc.add_paragraph()._p.append(omml_elem)
            with NamedTemporaryFile(suffix=".docx", delete=False) as tmp_file:
                tmp_path = tmp_file.name
                tmp_doc.save(tmp_path)
            return pypandoc.convert_file(tmp_path, to="markdown").strip()
        except (OSError, RuntimeError) as e:
            _log.warning("omml_conversion_failed", extra={"error": str(e)})
            return "<--OMML conversion failed. See logs.-->"
        finally:
            if tmp_path:
                try:
                    os.unlink(tmp_path)
                except Exception:
                    pass

    def render_hyperlink(self, hyperlink_elem) -> str:
        label = hyperlink_elem.text
        rel_id = hyperlink_elem.get(f"{{{self.NAMESPACES['r']}}}id")
        if rel_id and rel_id in self._rels:
            return f"[{label}]({self._rels[rel_id]})"
        return label or ""

    def get_para_prefix(self, para: Paragraph, style_levels: dict) -> str:
        style_name = para.style.name or ""
        if style_name.startswith("Heading"):
            level = int(style_name.split()[-1])
            return f"{'#' * level} "
        ilvl = self.get_list_level_from_paragraph(para, style_levels)
        if ilvl is not None:
            return f"{'\t' * ilvl}- "
        return ""

    def render_para_runs(self, para: Paragraph) -> tuple[str, str]:
        parts: list[str] = []
        footnotes: list[str] = []

        w = self.NAMESPACES["w"]
        m = self.NAMESPACES["m"]

        for child in para._p.iterchildren():
            tag = child.tag
            if tag == f"{{{w}}}r":
                if child.text:
                    parts.append(child.text)
                for fn in child.iter(f"{{{w}}}footnoteReference"):
                    ref = fn.get(f"{{{w}}}id")
                    parts.append(f"[^{ref}]")
                    footnotes.append(f"[^{ref}] {self._footnotes.get(ref, '')}")
            elif tag == f"{{{w}}}hyperlink":
                parts.append(self.render_hyperlink(child))
            elif tag in (f"{{{m}}}oMath", f"{{{m}}}oMathPara"):
                parts.append(self.omml_to_md_latex(child))

        return "".join(parts), "\n".join(footnotes)

    def para_to_markdown(self, para: Paragraph, style_levels: dict) -> str:
        prefix = self.get_para_prefix(para, style_levels)
        body, footnotes = self.render_para_runs(para)
        if not body.strip():
            return ""
        if footnotes:
            return f"{prefix}{body}\n{footnotes}".rstrip()
        return f"{prefix}{body}".rstrip()

    def format_cell_text(self, cell_text: list[str]) -> str:
        def normalize(t: str) -> str:
            return t.replace("\n", "<br>").replace("\t", "&nbsp; ")
        return "<br>".join(normalize(t) for t in cell_text if t).strip()

    def extract_table_text(self, table: Table) -> list[list[str]]:
        rows = []
        for row in table.rows:
            row_cells = []
            for cell in row.cells:
                cell_text = [self.para_to_markdown(p, self._style_levels) for p in cell.paragraphs]
                row_cells.append(self.format_cell_text(cell_text))
            rows.append(row_cells)
        return rows

    def table_to_markdown(self, table: Table) -> str:
        rows = self.extract_table_text(table)
        md_rows = ["| " + " | ".join(r) + " |" for r in rows]
        if rows:
            header_sep = "| " + " | ".join(["---"] * len(rows[0])) + " |"
            md_rows.insert(1, header_sep)
        return "\n".join(md_rows)

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def get_builtin_metadata(self, file: str) -> dict:
        metadata = {}
        with zipfile.ZipFile(file) as zf:
            root = etree.fromstring(zf.read("docProps/core.xml"))
            for elem in root:
                tag = etree.QName(elem).localname
                metadata[self.normalize_key(tag)] = self.clean_string(elem.text)

            if "docProps/custom.xml" in zf.namelist():
                root = etree.fromstring(zf.read("docProps/custom.xml"))
                for prop in root.findall("ep:property", self.NAMESPACES):
                    name = prop.get("name")
                    val_elem = next(iter(prop))
                    metadata[self.normalize_key(name)] = self.clean_string(val_elem.text)
        return metadata

    def get_metadata(self, builtIn_only: bool = True, mask_name: str | None = None) -> dict:
        stem = Path(self.filepath).stem
        builtin = self.get_builtin_metadata(self.filepath)
        metadata = {**builtin, "source": stem, "source_uri": str(self.filepath)}
        if mask_name:
            metadata = MetadataFilter().filter(metadata, mask_name)
        self.meta = metadata
        return metadata

    def to_markdown(self) -> str:
        lines = []
        for block in self.doc.iter_inner_content():
            if isinstance(block, Paragraph):
                md = self.para_to_markdown(block, self._style_levels)
                if md:
                    lines.append(md)
            elif isinstance(block, Table):
                md = self.table_to_markdown(block)
                if md:
                    lines.append(md)
        self.body = "\n\n".join(lines)
        _log.debug("docx_to_markdown", extra={"file": self.filepath, "n_blocks": len(lines)})
        return self.body


class ITRDOCXParser(DOCXParser):
    """DOCXParser extended for iTrust document templates (cover-page metadata)."""

    def __init__(self, file=None):
        super().__init__(file=file)

    def get_metadata_from_cover_page(self, file: str) -> dict:
        doc = Document(file)
        if len(doc.tables) < 2:
            return {}
        rows = self.extract_table_text(doc.tables[1])
        metadata = {}
        for row in rows:
            if len(row) == 2 and row[0].strip():
                metadata[row[0].strip()] = row[1].strip()
        return metadata

    def get_metadata(self, builtIn_only: bool = True, mask_name: str | None = None) -> dict:
        stem = Path(self.filepath).stem
        builtin = self.get_builtin_metadata(self.filepath)
        metadata = {**builtin, "source": stem, "source_uri": str(self.filepath)}

        if not builtIn_only:
            try:
                custom = self.get_metadata_from_cover_page(self.filepath)
                metadata = {**metadata, **custom}
            except (IndexError, KeyError) as e:
                _log.warning("cover_page_metadata_failed", extra={"file": self.filepath, "error": str(e)})

        if mask_name:
            metadata = MetadataFilter().filter(metadata, mask_name)

        self.meta = metadata
        return metadata
