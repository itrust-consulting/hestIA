from __future__ import annotations

import logging
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt

from hestia.domain.exceptions import ValidationError
from hestia.infrastructure.parsers.base import BaseParser, MetadataFilter

_log = logging.getLogger("hestia.system")


class PPTXParser(BaseParser):

    def __init__(self, file=None):
        self._loader_map = {".pptx": self._load_pptx}
        super().__init__(file=file)

    def _load_pptx(self, file) -> Presentation:
        prs = Presentation(file)
        _log.debug("pptx_loaded", extra={"file": str(file), "n_slides": len(prs.slides)})
        return prs

    # ------------------------------------------------------------------
    # Metadata
    # ------------------------------------------------------------------

    def get_metadata(self, builtIn_only: bool = True, mask_name: str | None = None) -> dict:
        stem = Path(self.filepath).stem
        metadata: dict = {"source": stem, "source_uri": str(self.filepath)}

        if self.doc is not None:
            props = self.doc.core_properties
            for attr in ("title", "subject", "author", "description", "version", "language"):
                val = getattr(props, attr, None)
                if val:
                    metadata[self.normalize_key(attr)] = self.clean_string(str(val))
            modified = getattr(props, "modified", None)
            if modified:
                metadata["modified"] = modified.isoformat()

        if mask_name:
            metadata = MetadataFilter().filter(metadata, mask_name)

        self.meta = metadata
        _log.debug("pptx_get_metadata", extra={"file": self.filepath, "keys": list(metadata.keys()), "mask": mask_name})
        return metadata

    # ------------------------------------------------------------------
    # Conversion
    # ------------------------------------------------------------------

    def to_markdown(self) -> str:
        if self.doc is None:
            raise ValidationError("Empty document. Parse a file first.")

        sections: list[str] = []
        for idx, slide in enumerate(self.doc.slides, start=1):
            title = self._slide_title(slide)
            heading = f"## Slide {idx}: {title}" if title else f"## Slide {idx}"
            lines: list[str] = [heading]
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                shape_title = shape.name
                # Skip the title placeholder — already in the heading
                if shape.shape_id == slide.shapes.title.shape_id if slide.shapes.title else False:
                    continue
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if not text:
                        continue
                    level = para.level
                    prefix = ("  " * level + "- ") if level > 0 else "- "
                    lines.append(f"{prefix}{text}")
            # Only include slides that have content beyond the heading
            if len(lines) > 1:
                sections.append("\n".join(lines))

        self.body = "\n\n".join(sections)
        _log.debug("pptx_to_markdown", extra={"file": self.filepath, "n_slides": len(self.doc.slides)})
        return self.body

    def _slide_title(self, slide) -> str:
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            return self.clean_string(slide.shapes.title.text)
        return ""
